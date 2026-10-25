"""Build the report from the completed, audited simulation artifacts."""
import json
from pathlib import Path
from statistics import mean
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
data = json.loads((ROOT / "research/starter-sid-recovery-simulation.json").read_text(encoding="utf-8"))
audit = json.loads((ROOT / "research/starter-sid-stall-audit.json").read_text(encoding="utf-8"))
rows, summary, config = data["rows"], data["summary"], data["config"]
overall, by = summary["overall"], summary["byStarter"]
names = ["Treecko", "Torchic", "Mudkip"]
fallback = {(r["trial"],r["species"]):r for r in audit["rows"]}
combined = {n:mean(r["experimentCount"]+fallback.get((r["trial"],n),{}).get("additionalCheckpoints",0) for r in rows if r["species"]==n) for n in names}
manual_steps = [s for r in audit["rows"] for s in r["steps"]]
dest = ROOT / "artifacts/Emerald-SID-Recovery-Simulation-Report.docx"
dest.parent.mkdir(exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.6)
sec.left_margin, sec.right_margin = Inches(.7), Inches(.7)
sec.header_distance, sec.footer_distance = Inches(.28), Inches(.28)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"]:
    style = doc.styles[name]
    style.font.name = "Aptos"
    style.font.color.rgb = RGBColor(0, 0, 0)
    if style.element.rPr is not None:
        for color in style.element.rPr.findall(qn("w:color")):
            for attr in ["themeColor", "themeTint", "themeShade"]:
                color.attrib.pop(qn("w:" + attr), None)
for style in doc.styles:
    for border in style.element.findall(".//" + qn("w:pBdr")):
        border.getparent().remove(border)
normal = doc.styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
doc.styles["Title"].font.size = Pt(25)
doc.styles["Title"].paragraph_format.space_after = Pt(10)
doc.styles["Heading 1"].font.size = Pt(17)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(4)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(10)
doc.styles["Heading 2"].font.size = Pt(12)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
for s in ["Heading 1", "Heading 2"]:
    doc.styles[s].paragraph_format.keep_with_next = True
header = sec.header.paragraphs[0]
header.text = "EMERALD CURTAIN INVESTIGATOR  |  SIMULATION REPORT"
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("24 October 2026  ·  ").font.size = Pt(8)
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = "Emerald SID recovery simulation"
doc.core_properties.subject = "Blind simulation of the guided starter route and curtain experiments"
doc.core_properties.author = "Emerald Curtain Investigator"

def p(text, bold=False, small=False):
    para = doc.add_paragraph()
    run = para.add_run(text); run.bold = bold
    if small: run.font.size = Pt(9)
    return para

def h(text, level=1):
    return doc.add_heading(text, level)

def page(title):
    doc.add_page_break(); h(title)

def table(headers, records, widths, align=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    tblpr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement("w:" + edge)
        for attr, value in [("val", "single"), ("sz", "5"), ("color", "D9D9D9")]: el.set(qn("w:" + attr), value)
        borders.append(el)
    tblpr.append(borders)
    for i, width in enumerate(widths): t.columns[i].width = Inches(width)
    repeat = OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, text in enumerate(headers): t.rows[0].cells[i].text = str(text)
    for record in records:
        cells = t.add_row().cells
        for i, value in enumerate(record): cells[i].text = str(value)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + edge); el.set(qn("w:w"), "85"); el.set(qn("w:type"), "dxa"); margins.append(el)
            tcpr.append(margins)
            if ri == 0 or ri % 2 == 0:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "243C52" if ri == 0 else "F2F5F8"); tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if (align or [0])[ci if align else 0] == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(10)
                    if ri == 0:
                        run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def link(label, url):
    para = doc.add_paragraph()
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "245D4A"); props.append(color)
    run.append(props); text = OxmlElement("w:t"); text.text = label; run.append(text); hyperlink.append(run); para._p.append(hyperlink)
    para.paragraph_format.space_after = Pt(4)

def case_name(r): return f"{r['trial']:03d} {r['species']}"
def ivs(r): return " / ".join(map(str, r["starter"]["ivs"]))
def decimal(v): return f"{v:.2f}"
successes = [r for r in rows if r["status"] == "recovered"]
stalled = [r for r in rows if r["status"] != "recovered"]
hardest = max(successes, key=lambda r: (r["experimentCount"], r["afterDrawing"], r["pixelQuestions"]))
easy = next(r for r in rows if r["status"] == "recovered" and r["experimentCount"] == 0)

doc.add_paragraph("Emerald SID recovery simulation", "Title")
p("100 paired scenarios for each starter", bold=True)
p("Prepared for the Emerald Curtain Investigator project · 24 October 2026", small=True)
p(f"The current guided workflow uniquely recovered the correct SID in {overall['recovered']} of {overall['runs']} runs. "
  f"It retained the true SID in every run, but {overall['unresolved']} runs stopped with multiple candidates because the guide offered no further experiment. "
  "The current workflow therefore did not achieve 100 out of 100 recoveries for every starter.")
p(f"Of the {overall['runs']} runs, {overall['zeroExperiments']} needed only the first complete curtain drawing. "
  f"Successful recoveries needed an average of {decimal(overall['meanExperimentsSuccessful'])} additional experiments, with a maximum of {overall['maxExperimentsSuccessful']}. "
  "These counts exclude route preparation and distinguish an experiment from answering one pixel.")
h("Recovery results", 2)
table(["Measure", *names], [
    ["Correct SID recovered", *[f"{by[n]['recovered']} / 100" for n in names]],
    ["Initial drawing was enough", *[by[n]['zeroExperiments'] for n in names]],
    ["Unresolved when guide stopped", *[by[n]['unresolved'] for n in names]],
    ["Mean experiments among recoveries", *[decimal(by[n]['meanExperimentsSuccessful']) for n in names]],
    ["Maximum among recoveries", *[by[n]['maxExperimentsSuccessful'] for n in names]],
    ["Recovered with diagnostic manual fallback", *[f"{by[n]['recovered']+sum(r['manuallyRecovered'] for r in audit['rows'] if r['species']==n)} / 100" for n in names]],
], [3.15,1.3,1.3,1.3], [0,1,1,1])
p(f"A separate diagnostic resolved {audit['manuallyRecovered']} of the {audit['primaryStalls']} stalled cases by allowing battle-only EV checkpoints. "
  "Those diagnostic recoveries are not counted as successes of the current guide.", bold=True)
p("This is a mathematical test of the existing workflow under accurate inputs and the stated timing model. It does not establish a real-world success rate or prove that the timing window is safe on every setup.", small=True)

page("How the simulation worked")
p("Each of 100 scenarios received a random public TID. All three starters were then tested against that same TID and SID, using the same sequence of possible starter timings. "
  "The first uniquely identifiable attempt could differ by species because their base stats differ. This produced 300 runs without treating them as 300 independent trainer scenarios.")
table(["Setting", "Value"], [
    ["TID and SID countdowns", "5 seconds, then 20 seconds"],
    ["SID generation timing", "Nominal advance 1245; estimated dialogue delay +50"],
    ["SID search scope", "1185–1305 inclusive, or ±60 advances"],
    ["Starter timing", "5 seconds, then 20 seconds; target advance 1195"],
    ["Starter search and actual timing error", "±60 advances; actual errors sampled uniformly from −60 to +60"],
    ["Calibration", "Zero for both timers"],
], [2.7,4.35], [0,0])
h("Keeping the answer hidden", 2)
p("An independent oracle generated the SID and Method 1 starter PID and IVs with BigInt arithmetic. It also packed a checksum-valid, encrypted 80-byte Pokémon record and converted each nibble to Black or Transparent. "
  "The planner and solver received no secret SID, exact IVs or PID from this oracle. Exact starter data became available only after the ordinary summary-stat search returned one spread.")
for text in [
    "Read all six level-5 stats, nature, gender and ability after Birch’s first Zigzagoon. Retry whenever the real starter search returns multiple PID and IV spreads.",
    "Generate surviving SIDs from the public TID and timing window. Ask the existing route planner to choose moves, EV training, held item and fallback items without consulting the true SID.",
    "Execute the checklist and record its EV and move changes. Draw all 160 curtain cells, then run the existing deduction engine.",
    "Follow the same experiment filters and ranking used by the guide. Answer only the useful requested pixels for held-item and move tests, reranking after each answer.",
    "Treat EV training as permanent. Create a fresh full drawing checkpoint and retain earlier observations. Stop at one SID or when the current guide has no available next experiment.",
]:
    para=doc.add_paragraph(text, "List Number"); para.paragraph_format.space_after=Pt(4)
p("One experiment means one held-item setup, move swap or EV checkpoint after the initial drawing. Collecting an item, choosing the first held item, individual pixel answers and starter retries are reported separately.", small=True)

page("Experiments and preparation")
p("Experiment averages below describe completed recoveries. A stalled run is not a quick success, so its partial experiment count is excluded from the recovery average. The CSV includes every run, including stalls.")
p(f"The initial timing filter left {min(r['initialSidCandidates'] for r in rows)}–{max(r['initialSidCandidates'] for r in rows)} distinct SIDs per run. "
  f"The first drawing reduced that to a mean of {decimal(overall['meanAfterDrawing'])}, with a maximum of {overall['maxAfterDrawing']}. "
  f"The current guide attempted {decimal(overall['meanExperimentsAttemptedAll'])} experiments per run overall, including the unfinished workloads of stalled runs.")
maximum=overall["maxExperimentsSuccessful"]
distribution=[]
for k in range(maximum+1):
    distribution.append([str(k), *[by[n]["successDistribution"].get(str(k),0) for n in names]])
distribution.append(["Unresolved", *[by[n]["unresolved"] for n in names]])
table(["Experiments to recover", *names], distribution, [3.15,1.3,1.3,1.3], [0,1,1,1])
h("Work before the first drawing", 2)
table(["Measure", *names], [
    ["Mean items listed for collection", *[decimal(by[n]["meanRouteItems"]) for n in names]],
    ["Mean distinct held items actually used", *[decimal(by[n]["meanDistinctItemsUsed"]) for n in names]],
    ["Runs with extra route EV training", *[by[n]["routeTraining"] for n in names]],
    ["Ambiguous first starter attempts", *[by[n]["ambiguousFirstAttempts"] for n in names]],
    ["Total starter retries", *[by[n]["starterRetries"] for n in names]],
], [3.15,1.3,1.3,1.3], [0,1,1,1])
p("The collection list covers potential branches across all surviving SIDs. A particular run may use few of those items because a transparent pixel or another experiment resolves it early. "
  "The gap between collected and used items suggests an opportunity for staged collection, although delaying collection could add return trips. This simulation did not optimize walking distance.")
p(f"Across all runs, the guide performed {overall['kinds']['item']} held-item experiments, {overall['kinds']['moves']} move swaps and {overall['kinds']['ev']} EV checkpoints. "
  f"It selected {overall['lastResortExperiments']} checkpoints marked as a last resort above 110 EVs. These are action counts, not elapsed-time measurements.", small=True)

page("Easy and difficult recoveries")
h("An initial drawing that was enough", 2)
p(f"Case {case_name(easy)} used TID {easy['tid']} and recovered SID {easy['trueSid']} immediately. The starter was at advance {easy['starter']['advance']}, "
  f"PID {easy['starter']['pid']}, with IVs {ivs(easy)} in HP, Attack, Defense, Speed, Sp. Attack, Sp. Defense order. "
  f"The initial {easy['initialSidCandidates']} SID candidates became one after drawing. The planned held item and move arrangement were already in place; this did not require zero preparation.")
h("The longest completed recovery", 2)
p(f"Case {case_name(hardest)} needed {hardest['experimentCount']} experiments after drawing. TID {hardest['tid']}, SID {hardest['trueSid']}, "
  f"starter advance {hardest['starter']['advance']}, PID {hardest['starter']['pid']}, {hardest['starter']['nature']} nature. "
  f"Its IVs were {ivs(hardest)}. The first drawing left {hardest['afterDrawing']} SID candidates.")
trace=[[e["number"],e["name"],e["before"],e["after"],"Redraw" if e.get("fullRedraw") else len(e["answers"])] for e in hardest["experiments"]]
table(["Test", "Action", "Before", "After", "Pixels"],trace,[.45,3.95,.75,.75,1.15],[1,0,1,1,1])
p("Before and After count surviving SIDs. Pixels counts explicitly answered questions; Redraw means the guide returned through a new EV checkpoint and the complete curtain was drawn again. "
  "Every recorded answer was generated from the actual encrypted record, and the true SID survived every reduction.",small=True)
h("What the starter comparison means", 2)
p("The table compares the current route and recommendation policy, including its availability gap. It does not prove a universal species ranking. "
  "Each starter’s move IDs, PP and unavoidable EVs differ, but its PID and IVs and the particular remaining SIDs determine whether those differences add useful evidence. "
  "A starter can require fewer tests among its successful runs while stalling more often overall.")
p("Torchic had the highest raw recovery count here at 80 out of 100, followed by Mudkip at 79 and Treecko at 74. "
  "Treecko’s broader early move options did not translate into the highest current-guide recovery rate in this sample. Fixing the recommendation gap and rerunning the comparison should precede a firm starter recommendation.",small=True)
p(f"The sample included {summary['distinctTids']} distinct TIDs, {summary['distinctSids']} distinct SIDs and " +
  ", ".join(f"{by[n]['distinctAcceptedPids']} accepted PIDs for {n}" for n in names) + ". "
  "Repeated starter PIDs are expected because every attempt uses the same narrow seed-zero timing window.",small=True)

page("Why the guide sometimes stopped")
p("The stalled cases retained the true SID, but the guide stopped offering tests. The diagnostic replay reproduced each stalled candidate set before trying any additional action, so these cases were not repaired by changing earlier observations.")
p("The core EV planner ranks vitamin-assisted and battle-only checkpoints together and removes duplicate digit tests in favour of cheaper plans. "
  "The beginner guide subsequently rejects every plan that requires vitamins. A cheaper vitamin plan can therefore remove a usable battle-only alternative before the availability filter runs. "
  "The result can be no suggested experiment even though a few ordinary battles would distinguish the remaining SIDs.")
table(["Diagnostic result", *names], [
    ["Stalled with current guide", *[by[n]["unresolved"] for n in names]],
    ["Resolved by manual EV continuation", *[sum(r["manuallyRecovered"] for r in audit["rows"] if r["species"]==n) for n in names]],
    ["Additional EV checkpoints", *[sum(r["additionalCheckpoints"] for r in audit["rows"] if r["species"]==n) for n in names]],
], [3.15,1.3,1.3,1.3], [0,1,1,1])
p(f"The diagnostic needed {audit['additionalCheckpoints']} additional checkpoints in total; {audit['lastResortCheckpoints']} were above 110 EVs. "
  "It chose checkpoints using the remaining candidate set and manual-training effort, then asked the hidden oracle for the resulting drawing. It did not use the correct SID to choose a target.")
p("Including the diagnostic fallback, the complete-process means were " + ", ".join(f"{decimal(combined[n])} experiments for {n}" for n in names) + ". "
  f"Manual fallback targets reached at most {max(s['target'] for s in manual_steps)} EVs; the largest single additional checkpoint required {max(s['battles'] for s in manual_steps)} one-EV battles. "
  "These figures describe the tested fallback policy, not a released fix or a promise of short play time.",small=True)
example=next((r for r in audit["rows"] if r["trial"]==9 and r["species"]=="Treecko"),audit["rows"][0])
original=next(r for r in rows if r["trial"]==example["trial"] and r["species"]==example["species"])
h("A concrete missed test",2)
p(f"Case {case_name(original)} stopped with SIDs " + " and ".join(f"{v} (0x{v:04X})" for v in original["finalSids"]) + ". "
  f"The diagnostic’s first action trained {example['steps'][0]['stat']} EVs to {example['steps'][0]['target']}, requiring {example['steps'][0]['battles']} one-EV battles. "
  "That checkpoint reduced two candidates to one.")
p("Recommended correction: apply the beginner route’s item and vitamin availability constraints before ranking and deduplicating EV plans. Keep a manual-training alternative whenever the faster vitamin version is unavailable. "
  "This report has not changed the app or counted that proposed correction as an implemented feature.",bold=True)

page("Limits and reproducibility")
h("What these results establish",2)
p("The simulation exercises the actual summary search, route planner, candidate solver, experiment ranking, pixel answers and accumulated evidence. An independent encoder checks the data permutation, XOR encryption and checksum, "
  "and a final exact solve verifies every run’s recorded evidence. No oracle SID is supplied as an identity fact or used to select a route or experiment.")
p("The negative control deliberately removed the true SID from a uniquely solved case’s allowed candidates. The solver then returned no SID. This checks that the harness cannot recover the hidden answer regardless of its evidence scope.")
h("What remains assumed",2)
p("SID creation occurs exactly at the model’s nominal advance 1245. This tests recovery inside the ±60 window, not whether real naming fades and dialogue delays always fall inside it. "
  "Starter input errors vary within the fixed window; timer settings and calibration do not vary. The first countdown is preparation time and is not added to the post-name SID interval.")
p("All stats, pixels, move order, item changes and EV logs are accurate. The route includes the avoidable two-Nincada trainer, uses one EXP recipient and cancels evolution. "
  "The oracle fixes unknown OT gender, met location and friendship to internally consistent values; they are not leaked to the planner. Extra EV-only opponents without a specified level use a level-5 nuisance EXP assumption. "
  "A cartridge playthrough is still needed to validate timing, route availability, encounter levels and the proposed recommendation fix. The simulator does not model missed pixels, lost battles, walking time, item budgets or the final shiny attempt.")
h("Files and commands",2)
p(f"Random seed: 0x{config['seed']:X}. Full inputs, routes, per-pixel reductions and source hashes are stored in starter-sid-recovery-simulation.json. "
  "The CSV has one row for each of the 300 runs; stalled counts are partial workloads, not experiments to successful recovery.",small=True)
p("Reproduce the primary study with: npx tsx scripts/simulate-starter-sid-recovery.ts 100. "
  "Then run npx tsx scripts/audit-starter-sid-stalls.ts for the separate diagnostic. The completed study was divided across eight independent processes; "
  "the seed and scenario generation are deterministic and results are merged by case and starter.",small=True)
h("Mechanics references",2)
commit="731ad5bfd6e6f265508d0efcca0ba42f9dcf5881"
link("Emerald source   Trainer ID and SID creation",f"https://github.com/pret/pokeemerald/blob/{commit}/src/new_game.c#L80")
link("Emerald source   Pokémon generation encryption and checksums",f"https://github.com/pret/pokeemerald/blob/{commit}/src/pokemon.c")
link("Emerald source   Stored moves EVs IVs and trainer data",f"https://github.com/pret/pokeemerald/blob/{commit}/include/pokemon.h")
link("PokéFinder   Generation 3 static PID and IV sequence","https://github.com/Admiral-Fish/PokeFinder/blob/master/Core/Gen3/Generators/StaticGenerator3.cpp")
p("The local engine source hashes in the JSON identify the exact implementation tested. No production logic or saved investigation was changed for this study.",small=True)

doc.save(dest)
print(dest)
